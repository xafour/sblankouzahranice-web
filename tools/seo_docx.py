#!/usr/bin/env python3
"""Z navrhy/SEO-NAVRHY.md udělá Wordový dokument ke schválení: u každého návrhu políčko „Souhlasím“ a sloupec na vlastní text.

Použití: python3 tools/seo_docx.py [navrhy/SEO-NAVRHY.md] [navrhy/SEO-NAVRHY-ke-schvaleni.docx]
Potřebuje balíček python-docx.
"""
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

SRC = sys.argv[1] if len(sys.argv) > 1 else "navrhy/SEO-NAVRHY.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else "navrhy/SEO-NAVRHY-ke-schvaleni.docx"
ZELENA, OLIVOVA, SEDA = "145A38", "DEE5A3", "F4F4EF"


def cisty(s):
    """Markdown → prostý text (odkazy, tučné, kurzíva, kód, počet znaků v závorce)."""
    s = re.sub(r"\*\(\d+ zn\.\)\*", "", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    s = s.replace("\\|", "|").replace("**", "").replace("`", "")
    return re.sub(r"(?<!\w)\*([^*]+)\*(?!\w)", r"\1", s).strip()


def radky_tabulky(blok):
    out = []
    for r in blok:
        bunky = [b.strip() for b in re.split(r"(?<!\\)\|", r.strip())[1:-1]]
        if bunky and not set("".join(bunky)) <= set("-: "):
            out.append(bunky)
    return out[1:]  # bez záhlaví


def nacti(md):
    """Rozdělí SEO-NAVRHY.md na: titulky+popisy, jen popisy, skrýt, obrázky, k ověření."""
    data = {"titulky": [], "popisy": [], "skryt": [], "obrazky": [], "overeni": []}
    kapitola, podkapitola, tabulka = None, None, []

    def zavri_tabulku():
        nonlocal tabulka
        for b in radky_tabulky(tabulka):
            if kapitola == "1":
                data["titulky"].append((cisty(b[0]), cisty(b[1]), cisty(b[2])))
            elif kapitola == "2":
                data["popisy"].append((cisty(b[0]), cisty(b[1])))
            elif kapitola == "3":
                data["obrazky"].append((podkapitola, cisty(b[0]), cisty(b[1])))
        tabulka = []

    for r in md.splitlines():
        if r.startswith("|"):
            tabulka.append(r)
            continue
        if tabulka:
            zavri_tabulku()
        if r.startswith("## "):
            kapitola = r[3] if r[3].isdigit() else ("over" if "ověření" in r else None)
        elif r.startswith("### "):
            podkapitola = cisty(r[4:])
        elif r.startswith("- ") and kapitola in ("1", "2") and ("Skrýt" in r or "skryt" in r or "skrýt" in r):
            data["skryt"].append(cisty(r[2:]).rstrip(",."))
        elif kapitola == "3" and r.startswith("Stránky:") and podkapitola and podkapitola.startswith("Boční"):
            podkapitola = "Boční panel (na 8 stránkách)"
        elif kapitola == "3" and r.startswith("**9×"):
            data["overeni"].append(f"{podkapitola}: {cisty(r)}")
        elif kapitola == "over" and r.startswith("- ") and not cisty(r).endswith(":"):
            data["overeni"].append(cisty(r[2:]))
        elif kapitola == "over" and r.startswith("  - "):
            data["overeni"].append(cisty(r[4:]))
        elif kapitola == "over" and r.startswith("  ") and r.strip() and data["overeni"]:
            data["overeni"][-1] += " " + cisty(r)  # pokračování položky na dalším řádku
    if tabulka:
        zavri_tabulku()
    return data


# --- Word ---------------------------------------------------------------

def podbarvi(bunka, barva):
    bunka._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{barva}"/>'))


def zaskrtavatko(bunka):
    """Klikací políčko ☐/☒ (ovládací prvek Wordu, funguje i v LibreOffice a Google Dokumentech jako znak)."""
    p = bunka.paragraphs[0]
    p.alignment = 1
    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")} xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">'
        '<w:sdtPr><w14:checkbox><w14:checked w14:val="0"/>'
        '<w14:checkedState w14:val="2612" w14:font="MS Gothic"/>'
        '<w14:uncheckedState w14:val="2610" w14:font="MS Gothic"/></w14:checkbox></w:sdtPr>'
        '<w:sdtContent><w:r><w:rPr><w:rFonts w:ascii="MS Gothic" w:eastAsia="MS Gothic" w:hAnsi="MS Gothic" w:hint="eastAsia"/>'
        '<w:sz w:val="32"/></w:rPr><w:t>☐</w:t></w:r></w:sdtContent></w:sdt>'
    )
    p._p.append(sdt)


def opakuj_zahlavi(radek):
    radek._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))


def tabulka(doc, zahlavi, sirky, radky, souhlas=True):
    """radky: seznam seznamů textů (bez sloupců Souhlasím / Vaše změna – ty se přidají)."""
    cele = zahlavi + (["Souhlasím", "Vaše změna (napište nový text)"] if souhlas else ["Vaše odpověď"])
    sirky = sirky + ([2.6, 7.1] if souhlas else [9.7])
    t = doc.add_table(rows=1, cols=len(cele))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, sirky):
        col.width = Cm(w)  # LibreOffice a Google Dokumenty čtou šířky z mřížky tabulky
    for i, (txt, w) in enumerate(zip(cele, sirky)):
        c = t.rows[0].cells[i]
        c.width = Cm(w)
        run = c.paragraphs[0].add_run(txt)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        podbarvi(c, ZELENA)
    opakuj_zahlavi(t.rows[0])
    for n, r in enumerate(radky):
        cells = t.add_row().cells
        for i, w in enumerate(sirky):
            cells[i].width = Cm(w)
        for i, txt in enumerate(r):
            if isinstance(txt, tuple):  # (text, poznámka šedě)
                cells[i].paragraphs[0].add_run(txt[0])
                pozn = cells[i].paragraphs[0].add_run(f"  {txt[1]}")
                pozn.font.size = Pt(8)
                pozn.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            else:
                cells[i].paragraphs[0].add_run(txt)
        if souhlas:
            zaskrtavatko(cells[len(r)])
        if n % 2:
            for c in cells[: len(r)]:
                podbarvi(c, SEDA)
    doc.add_paragraph()
    return t


def nadpis(doc, text, uroven=1):
    h = doc.add_heading(text, uroven)
    for r in h.runs:
        r.font.color.rgb = RGBColor.from_string(ZELENA)


def odstavec(doc, text, tucne=None):
    p = doc.add_paragraph()
    if tucne:
        p.add_run(tucne).bold = True
    p.add_run(text)
    return p


def main():
    d = nacti(open(SRC, encoding="utf-8").read())
    doc = Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21)
    for okraj in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, okraj, Cm(1.5))
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    nadpis(doc, "Návrhy textů pro Google – ke schválení", 0)
    odstavec(doc, "Připraveno 15. 10. 2026. Zatím jsou to jen návrhy, na webu se nic nezměnilo.")
    nadpis(doc, "Jak dokument vyplnit")
    for t, txt in [
        ("Souhlasíte s návrhem: ", "klikněte na políčko ☐ ve sloupci Souhlasím (změní se na ☒)."),
        ("Chcete jiný text: ", "napište ho do sloupce Vaše změna. Políčko pak nechte prázdné."),
        ("Návrh nechcete vůbec: ", "napište do sloupce Vaše změna „NE“."),
        ("Hotovo: ", "dokument uložte a pošlete zpět. Vše pak doplním na zkušební web hromadně, ručně nemusíte nic vyplňovat."),
    ]:
        odstavec(doc, txt, t).style = "List Bullet"
    odstavec(doc, "SEO titulek je modrý nadpis ve výsledcích Googlu (nejvýš 60 znaků). Meta popis jsou dva řádky pod ním "
             "(120–155 znaků). Popisek obrázku (alt) řekne Googlu a nevidomým, co je na fotce. Počet znaků je u návrhu šedě. "
             "Texty se vyhýbají slovům „průvodkyně“, „průvodce“ (o osobě) a „průvodcování“; jako průvodce jsou zmíněné jen "
             "vaše produkty (e-průvodci, průvodce do kapsy).")

    nadpis(doc, "1. Titulky a popisy stránek")
    odstavec(doc, "Tyto stránky mají dnes v Googlu stejný titulek jako úvodní stránka, takže je Google od sebe nerozezná. "
             "Titulek a popis úvodní stránky napíšete sama.")
    radky = []
    for stranka, tit, pop in d["titulky"]:
        radky.append([stranka, "SEO titulek", (tit, f"({len(tit)} zn.)")])
        radky.append(["", "Meta popis", (pop, f"({len(pop)} zn.)")])
    tabulka(doc, ["Stránka", "Co", "Můj návrh"], [4.2, 2.2, 10.6], radky)

    nadpis(doc, "2. Stránky bez popisu pro Google")
    tabulka(doc, ["Stránka", "Návrh meta popisu"], [5.0, 12.0],
            [[st_, (p, f"({len(p)} zn.)")] for st_, p in d["popisy"]])

    nadpis(doc, "3. Stránky, které navrhuji skrýt z Googlu")
    odstavec(doc, "Tyto stránky do výsledků Googlu nepatří (šablony, kopie, technické stránky).")
    radky = []
    for r in d["skryt"]:
        stranka, _, popis = r.partition(":")
        radky.append([stranka.strip(), popis.strip()])
    tabulka(doc, ["Stránka", "Co to je a co navrhuji"], [5.0, 12.0], radky)

    nadpis(doc, "4. Popisky obrázků (alt)")
    odstavec(doc, "Je to 57 různých fotek. Popisek se vyplní jednou u fotky a platí na všech stránkách, kde je použitá.")
    radky, minula = [], None
    for stranka, soubor, alt in d["obrazky"]:
        radky.append([stranka if stranka != minula else "", soubor, alt])
        minula = stranka
    tabulka(doc, ["Stránka", "Obrázek (soubor)", "Návrh popisku"], [4.0, 5.2, 7.8], radky)

    nadpis(doc, "5. Otázky k ověření")
    tabulka(doc, ["Otázka"], [17.0], [[q] for q in d["overeni"]], souhlas=False)

    doc.save(OUT)
    print(OUT, {k: len(v) for k, v in d.items()})


if __name__ == "__main__":
    main()
